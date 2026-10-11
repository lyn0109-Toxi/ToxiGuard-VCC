"""Calculation, review lifecycle, bilingual UI and exported-formula regressions."""
from __future__ import annotations

import ast
from copy import deepcopy
from io import BytesIO
from math import isclose
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
import pandas as pd
from streamlit.testing.v1 import AppTest

from analytical_workspace import (
    BUILD, CATALOG, MASS_MG, add_custom_test, blank_preparation, calculate_preparation,
    daily_amounts, demo_workspace, export_workbook, impurity_limit, new_workspace,
    record_review, refresh_reviews, report_markdown, select_tests,
)
from validation_extension import evaluate_q14_problem, q14_evidence_fingerprint, q14_problem_frame

def near(actual: float, expected: float) -> None:
    assert isclose(actual, expected, rel_tol=1e-10, abs_tol=1e-12), (actual, expected)

def rejects(action) -> None:
    try:
        action()
    except (ValueError, OverflowError):
        return
    raise AssertionError("Invalid input was accepted")

def validate_calculations() -> None:
    s = demo_workspace()
    d = daily_amounts(s["product"])
    near(d["api_mg_day"], 400)
    near(d["implied_api_mg_day"], 400)
    near(d["whole_g_day"], 1.2)
    record = s["tests"]["assay"]
    near(calculate_preparation(record["sample"], s["product"])["final_ug_ml"], 100)
    near(calculate_preparation(record["reference"], s["product"], True)["final_ug_ml"], 99.5)
    for unit, factor in MASS_MG.items():
        prep = deepcopy(record["reference"])
        prep.update(amount=20/factor, amount_unit=unit, water=90)
        near(calculate_preparation(prep, s["product"], True)["final_ug_ml"], 99.5)  # as-is never applies saved water
    prep.update(correction_mode="dry", potency=100, water=10)
    near(calculate_preparation(prep, s["product"], True)["final_ug_ml"], 90)
    prep["steps"].append({"aliquot_ml": 1, "final_ml": 10})
    near(calculate_preparation(prep, s["product"], True)["final_ug_ml"], 9)
    for field in ("amount", "stock_ml", "basis_factor"):
        for value in (None, 0, -1, float("nan"), float("inf"), True):
            invalid = deepcopy(prep)
            invalid[field] = value
            rejects(lambda: calculate_preparation(invalid, s["product"], True))
    for step in ({"aliquot_ml": 101, "final_ml": 200}, {"aliquot_ml": 5, "final_ml": 4}, {"aliquot_ml": 0, "final_ml": 10}):
        invalid = deepcopy(prep)
        invalid["steps"] = [step]
        rejects(lambda: calculate_preparation(invalid, s["product"], True))
    liquid = deepcopy(s["product"])
    liquid.update(strength=10, whole_unit=5, whole_unit_unit="mL", units_day=4, max_daily_api=40, density=None)
    assert daily_amounts(liquid)["whole_g_day"] is None
    liquid["density"] = 1.1
    near(daily_amounts(liquid)["whole_g_day"], 22)
    liquid_prep = deepcopy(record["sample"])
    liquid_prep.update(basis="product_volume", amount=5, amount_unit="mL")
    near(calculate_preparation(liquid_prep, liquid)["final_ug_ml"], 10)
    for material in ("api", "excipient", "starting_material"):
        raw = deepcopy(s["product"])
        raw["material_type"] = material
        assert daily_amounts(raw)["whole_g_day"] is None
        rejects(lambda: calculate_preparation(record["sample"], raw))
    excessive = deepcopy(s["product"])
    excessive["strength"] = 301
    rejects(lambda: calculate_preparation(record["sample"], excessive))
    limit = s["limit"]
    limit.update(limit_kind="Fictional exposure allowance", allowance=5, source="Fixture only", version="v1", location="§1", scope="Fictional oral tablet")
    near(impurity_limit(s)["ug_g"], 5/1.2)
    limit["denominator"] = "api_daily"
    near(impurity_limit(s)["ug_g"], 12.5)
    s["product"]["max_daily_api"] = 399
    rejects(lambda: impurity_limit(s))
    s["product"]["max_daily_api"] = None
    rejects(lambda: impurity_limit(s))  # derived strength never silently replaces declared API dose
    s["product"].update(whole_unit=2500, units_day=4)
    limit["denominator"] = "whole_product"
    matrix = blank_preparation()
    matrix.update(basis="whole_matrix", amount=.2, amount_unit="g", stock_ml=50, steps=[{"aliquot_ml":5,"final_ml":10}])
    near(impurity_limit(s, matrix)["solution_ug_ml"], .001)  # 1 ng/mL
    matrix["correction_mode"] = "as_is"
    matrix["potency"] = 99
    rejects(lambda: impurity_limit(s, matrix))
    s["limit"]["source"] = ""
    rejects(lambda: impurity_limit(s))

def validate_reviews() -> None:
    s = demo_workspace()
    r = s["tests"]["assay"]
    refresh_reviews(s)
    record_review(s, "assay", "sample", "Independent reviewer", "Recomputed with original inputs")
    record_review(s, "assay", "reference", "Independent reviewer", "Checked potency basis")
    r["sample"]["amount"] = 600
    refresh_reviews(s)
    assert r["sample"]["review"] is None and r["reference"]["review"]
    r["sample"]["amount"] = 300
    refresh_reviews(s)
    assert r["sample"]["review"] is None  # restoring values does not restore approval
    r["procedure"] = "New method version"
    refresh_reviews(s)
    assert r["reference"]["review"] is None
    rejects(lambda: record_review(s, "assay", "evidence", "Reviewer", "Only selected"))
    r.update(applicability="applicable", rationale="Fixture scope", source="Fixture monograph", version="v1", location="§1", scope="Fixture laboratory", excerpt="Fixture text", specification="Fixture criterion", procedure="Fixture method", preparation="Fixture preparation", sst="Fixture SST", verification_plan="Fixture verification")
    record_review(s, "assay", "evidence", "Reviewer", "Checked fixture source")
    assert r["evidence_review"] and r["experiment_review"] is None
    rejects(lambda: record_review(s, "assay", "experiment", "Reviewer", "No experiment"))
    r.update(validation_plan="Fixture protocol v1", experimental_data="Fixture report v1")
    record_review(s, "assay", "experiment", "Reviewer", "Checked stated scope")
    refresh_reviews(s, {"batch":"new"})
    assert r["experiment_review"] is None
    key = add_custom_test(s, "Custom pH")
    assert add_custom_test(s, "Custom pH") == key and s["selected"].count(key) == 1
    select_tests(s, [])
    assert key in s["tests"]  # deselection preserves work
    q14 = q14_problem_frame().iloc[0].copy()
    q14["Status"] = "Defined"
    assert evaluate_q14_problem(q14) == "Review"
    for field in ("Source document", "Source version", "Source location", "Application scope", "Evidence reviewer"):
        q14[field] = "Fixture evidence"
    q14["Evidence status"] = "Reviewer confirmed"
    q14["Evidence fingerprint"] = q14_evidence_fingerprint(q14)
    assert evaluate_q14_problem(q14) == "Pass"
    q14["Source version"] = "v2"
    assert evaluate_q14_problem(q14) == "Review"

def formula_value(wb, sheet: str, address: str):
    """Evaluate only the emitted Excel subset, independently from app calculation."""
    value = wb[sheet][address].value
    if not isinstance(value, str) or not value.startswith("="):
        return value
    expression = value[1:]
    expression = re.sub(r"'([^']+)'!([A-Z]+[0-9]+)", lambda m: f'cell("{m[1]}","{m[2]}")', expression)
    expression = re.sub(r'(?<![!"A-Z0-9])([A-Z]+[0-9]+)(?!["A-Z0-9])', lambda m: f'cell("{sheet}","{m[1]}")', expression)
    expression = re.sub(r"(?<![<>=!])=(?!=)", "==", expression)
    def evaluate(node):
        if isinstance(node, ast.Constant): return node.value
        if isinstance(node, ast.BinOp):
            a, b = evaluate(node.left), evaluate(node.right)
            if isinstance(node.op, ast.Add): return a+b
            if isinstance(node.op, ast.Sub): return a-b
            if isinstance(node.op, ast.Mult): return a*b
            if isinstance(node.op, ast.Div): return a/b
        if isinstance(node, ast.Compare) and isinstance(node.ops[0], ast.Eq): return evaluate(node.left) == evaluate(node.comparators[0])
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
            if node.func.id == "IF": return evaluate(node.args[1] if evaluate(node.args[0]) else node.args[2])
            if node.func.id == "cell": return formula_value(wb, evaluate(node.args[0]), evaluate(node.args[1]))
        raise AssertionError(ast.dump(node))
    return evaluate(ast.parse(expression, mode="eval").body)

def validate_exports() -> None:
    s = demo_workspace()
    s["product"]["name"] = '=HYPERLINK("https://example.invalid","not executable")'
    for lang in ("en", "ko"):
        report = report_markdown(s, lang)
        assert ("분석법 작업공간" if lang == "ko" else "Analytical method workspace") in report
        assert "99.5 ug/mL" in report and "100 ug/mL" in report
        assert ("미검증" if lang == "ko" else "unvalidated") in report
        wb = load_workbook(BytesIO(export_workbook(s, lang)), data_only=False)
        assert wb["Product"]["B3"].data_type == "s"
        near(formula_value(wb, "02_sample", "B17"), 100)
        near(formula_value(wb, "02_reference", "B17"), 99.5)
        near(formula_value(wb, "02_sample", "B19"), 0)
        wb["02_sample"]["B2"] = .3
        wb["02_sample"]["C2"] = "g"
        near(formula_value(wb, "02_sample", "B17"), 100)
        wb["02_reference"]["B6"] = 100
        near(formula_value(wb, "02_reference", "B17"), 100)
        assert wb.calculation.fullCalcOnLoad
        # Blank examples have no output formulas that look like valid calculations.
        blank_sheet = wb["01_sample"]
        assert not any(c.data_type == "f" for row in blank_sheet for c in row)
    s["limit"].update(limit_kind="Fictional", allowance=5, source="Fixture", version="v1", location="§1", scope="Oral fixture")
    wb = load_workbook(BytesIO(export_workbook(s)), data_only=False)
    near(formula_value(wb, "Exposure", "B10"), 5/1.2)
    s["limit"]["denominator"] = "api_daily"
    wb = load_workbook(BytesIO(export_workbook(s)), data_only=False)
    near(formula_value(wb, "Exposure", "B10"), 12.5)

def validate_ui() -> None:
    page = AppTest.from_file(str(ROOT / "streamlit_app.py"))
    page.run(timeout=30)
    page.button(key="landing_method_button").click().run(timeout=30)
    assert not page.exception and page.session_state["active_page"] == "method"
    assert not page.session_state["analytical_method_workspace"]["selected"]
    page.button(key="am_demo").click().run(timeout=30)
    page.selectbox(key="am_active_method_test").set_value("assay").run(timeout=30)
    page.selectbox(key="am_active_calculation_test").set_value("assay").run(timeout=30)
    assert not page.exception
    values = [v.value for v in page.metric]
    assert "100 ug/mL" in values and "99.5 ug/mL" in values
    for kind in ("sample", "reference"):
        page.text_input(key="am_review_assay_"+kind+"_reviewer").set_value("Fixture reviewer")
        page.text_input(key="am_review_assay_"+kind+"_note").set_value("Independent recomputation")
        page.button(key="am_review_assay_"+kind+"_button").click().run(timeout=30)
    r = page.session_state["analytical_method_workspace"]["tests"]["assay"]
    assert r["sample"]["review"] and r["reference"]["review"]
    page.number_input(key="am_tests_assay_sample_amount").set_value(600).run(timeout=30)
    r = page.session_state["analytical_method_workspace"]["tests"]["assay"]
    assert r["sample"]["review"] is None and r["reference"]["review"]
    page.number_input(key="am_tests_assay_sample_amount").set_value(300).run(timeout=30)
    assert page.session_state["analytical_method_workspace"]["tests"]["assay"]["sample"]["review"] is None
    page.selectbox(key="am_active_calculation_test").set_value("identity").run(timeout=30)
    page.selectbox(key="am_active_calculation_test").set_value("assay").run(timeout=30)
    assert page.number_input(key="am_tests_assay_sample_amount").value == 300
    page.radio(key="language_choice").set_value("ko").run(timeout=30)
    assert not page.exception
    assert page.selectbox(key="am_active_calculation_test").value == "assay"
    assert page.number_input(key="am_tests_assay_sample_amount").label == "분석 시료 / 표준품 채취량"
    assert page.selectbox(key="am_active_calculation_test").options[1] == "함량 / 역가"
    page.text_input(key="am_custom_name").set_value("Custom pH").run(timeout=30)
    page.button(key="am_add_custom").click().run(timeout=30)
    page.button(key="am_add_custom").click().run(timeout=30)
    assert page.session_state["analytical_method_workspace"]["selected"].count(next(k for k in page.session_state["analytical_method_workspace"]["tests"] if k.startswith("custom_"))) == 1
    page.button(key="workflow_response").click().run(timeout=30)
    assert not page.exception
    assert BUILD in page.text_area[-1].value
    page.text_input(key="profile_batch").set_value("Changed batch").run(timeout=30)
    assert page.session_state["analytical_method_workspace"]["tests"]["assay"]["reference"]["review"] is None
    page.button(key="workflow_method").click().run(timeout=30)
    page.selectbox(key="am_product_material_type").set_value("api").run(timeout=30)
    assert not page.exception
    assert daily_amounts(page.session_state["analytical_method_workspace"]["product"])["whole_g_day"] is None
    page.selectbox(key="am_product_material_type").set_value("finished_product").run(timeout=30)
    assert page.number_input(key="am_product_whole_unit").value == 300
    # Partially filled dynamic dilution rows stay incomplete without NaN hash crashes.
    page.session_state["am_steps_assay_sample_ko"] = {"edited_rows": {}, "added_rows": [{"분취량 (mL)": 1}], "deleted_rows": []}
    page.run(timeout=30)
    assert not page.exception
    page.radio(key="language_choice").set_value("en").run(timeout=30)
    assert not page.exception
    assert page.number_input(key="am_tests_assay_reference_potency").label == "Assigned purity / potency (%)"

if __name__ == "__main__":
    for check in (validate_calculations, validate_reviews, validate_exports, validate_ui):
        check()
        print(check.__name__, "PASS")
