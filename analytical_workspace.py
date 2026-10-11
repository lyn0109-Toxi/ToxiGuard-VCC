"""Explicit analytical-method records and dimensional calculations (no preset limits)."""
from __future__ import annotations

from copy import deepcopy
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
from math import isclose, isfinite
from typing import Any

BUILD = "analytical-workspace-2026-10-11-v1"
MASS_MG = {"g": 1000.0, "mg": 1.0, "ug": 0.001, "ng": 0.000001}
VOLUME_ML = {"mL": 1.0, "L": 1000.0}
CATALOG = {
    "identity": ("Identity", "확인", "Confirm material identity", "물질의 동일성 확인"),
    "assay": ("Assay / potency", "함량 / 역가", "Measure content or potency", "함량 또는 역가 측정"),
    "related_substances": ("Related substances", "유연물질", "Assess process impurities and degradation products", "공정불순물 및 분해산물 평가"),
    "residual_solvents": ("Residual solvents", "잔류용매", "Assess process solvent residues", "공정 용매 잔류 평가"),
    "elemental_impurities": ("Elemental impurities", "원소불순물", "Assess elemental exposure and source contributions", "원소 노출 및 기여원 평가"),
    "mutagenic_impurities": ("Mutagenic impurities", "변이원성 불순물", "Assess DNA-reactive impurity controls", "DNA 반응성 불순물 관리 평가"),
    "nitrosamines": ("Nitrosamines", "니트로사민", "Assess product-specific nitrosamine risk", "제품별 니트로사민 위해 평가"),
    "water": ("Water / loss on drying", "수분 / 건조감량", "Assess water or volatile loss on the stated basis", "명시한 기준의 수분 또는 휘발 손실 평가"),
    "dissolution": ("Dissolution", "용출", "Assess release performance", "방출 성능 평가"),
    "disintegration": ("Disintegration", "붕해", "Assess dosage-form disintegration", "제형의 붕해 평가"),
    "uniformity": ("Dosage-unit uniformity", "제제균일성", "Assess uniformity on the justified basis", "근거가 있는 기준으로 균일성 평가"),
    "particle_size": ("Particle size", "입도", "Assess performance-related size distribution", "성능 관련 입도분포 평가"),
    "solid_form": ("Solid form", "결정형 / 고체상", "Assess relevant crystalline or solid-state form", "관련 결정형 또는 고체상 평가"),
    "microbiology": ("Microbiological quality", "미생물", "Assess relevant microbiological quality", "관련 미생물 품질 평가"),
    "sterility": ("Sterility", "무균", "Assess sterility test applicability and suitability", "무균시험 적용성 및 적합성 평가"),
    "endotoxin": ("Bacterial endotoxins", "엔도톡신", "Assess route-specific endotoxin controls", "경로별 엔도톡신 관리 평가"),
    "appearance": ("Appearance", "성상", "Describe relevant physical attributes", "관련 물리적 특성 기술"),
    "other": ("Other / product-specific tests", "기타 / 제품별 시험", "Record additional justified tests", "추가 시험의 근거 기록"),
}

# Labels are also the report/export contract. Source wording is never translated.
LABELS = {
    "active_method_test": ("Method record", "시험법 기록"), "active_calculation_test": ("Concentration test", "농도 계산 시험"),
    "field": ("Input field", "입력 항목"), "custom_name": ("Custom test name", "사용자 시험항목명"),
    "material_type": ("Material / product type", "원료 / 제품 구분"),
    "name": ("Material / product name", "원료 / 제품명"), "api_name": ("API name", "API명"),
    "identifier": ("Identifier / CAS", "식별자 / CAS"), "form": ("Salt / hydrate / stereochemical form", "염 / 수화물 / 입체형"),
    "claim_basis": ("Label claim / analyte basis", "표시량 / 분석대상 기준"),
    "dosage_form": ("Dosage form", "제형"), "route": ("Route", "투여경로"), "market": ("Target market", "대상 시장"),
    "manufacturer": ("Manufacturer / matrix", "제조원 / 매트릭스"), "batch": ("Batch", "배치"), "stage": ("Intended use / stage", "의도한 사용 / 단계"),
    "max_daily_api": ("Maximum daily API dose", "최대 일일 API 투여량"), "api_dose_unit": ("API dose unit / day", "API 투여량 단위 / 일"),
    "strength": ("API strength per dosage unit", "투여단위당 API 표시량"), "strength_unit": ("Strength mass unit", "표시량 질량 단위"),
    "whole_unit": ("Whole product amount per dosage unit", "투여단위당 전체 제품량"), "whole_unit_unit": ("Whole product mass / volume unit", "전체 제품 질량 / 부피 단위"),
    "units_day": ("Maximum dosage units / day", "최대 투여단위수 / 일"), "density": ("Product density (g/mL); needed for volume-to-mass conversion", "제품 밀도 (g/mL); 부피에서 질량 환산 시 필요"),
    "max_daily_material": ("Daily material contribution (not an API dose)", "일일 원료 기여량 (API 투여량과 별도)"), "material_dose_unit": ("Material contribution unit / day", "원료 기여량 단위 / 일"),
    "purpose": ("Purpose", "시험 목적"), "applicability": ("Applicability", "적용성"), "rationale": ("Applicability / exclusion rationale", "적용 / 제외 근거"),
    "source_type": ("Source type", "출처 유형"), "source": ("Source document / identifier", "근거 문서 / 식별자"), "version": ("Version / effective date", "버전 / 시행일"), "location": ("Section / page / table / cell", "절 / 페이지 / 표 / 셀"),
    "scope": ("Applicable product / laboratory / method scope", "적용 제품 / 실험실 / 시험법 범위"), "excerpt": ("Original source excerpt", "원문 발췌"),
    "specification": ("Specification / acceptance criterion and basis", "규격 / 허용기준 및 근거"),
    "procedure": ("Procedure / conditions", "분석절차 / 조건"), "preparation": ("Sample / standard preparation", "검액 / 표준액 조제법"),
    "sst": ("SST / controls / failure actions", "시스템적합성 / 관리 / 실패 조치"),
    "verification_plan": ("Compendial applicability / verification plan", "공정서 적용 확인 / 검증 계획"),
    "validation_plan": ("Validation scope / protocol and criteria", "밸리데이션 범위 / 계획서 및 기준"), "experimental_data": ("Experimental report / version / results / limitations", "실험 보고서 / 버전 / 결과 / 한계"),
    "approach": ("Q14 development approach", "Q14 개발 접근법"), "value_kind": ("Value provenance", "값의 성격"),
    "amount": ("Analytical sample / standard amount", "분석 시료 / 표준품 채취량"), "amount_unit": ("Amount unit", "채취량 단위"),
    "basis": ("Amount represents", "채취량의 대상"), "correction_mode": ("Purity / potency basis", "순도 / 역가 기준"), "potency": ("Assigned purity / potency (%)", "지정 순도 / 역가 (%)"),
    "water": ("Water (%) on dry basis only", "수분 (%) 건조물 기준에만 적용"), "basis_factor": ("Analyte / salt conversion factor", "분석대상 / 염 환산계수"), "basis_note": ("Correction / conversion rationale", "보정 / 환산 근거"),
    "extraction": ("Assumed extraction fraction (%)", "가정한 추출 분율 (%)"), "stock_ml": ("Stock final volume (mL)", "원액 최종부피 (mL)"),
    "aliquot_ml": ("Aliquot (mL)", "분취량 (mL)"), "final_ml": ("Diluted final volume (mL)", "희석 후 최종부피 (mL)"),
    "reviewer": ("Reviewer", "검토자"), "review_note": ("Independent review note", "독립 검토 기록"),
    "pending": ("Unreviewed", "미검토"), "reviewed": ("Review recorded", "검토 기록됨"), "incomplete": ("Inputs incomplete", "입력 미완료"),
    "entered": ("User-entered", "사용자 입력"), "source_derived": ("Source-derived; verification separate", "출처 기반; 확인 상태 별도"), "calculated": ("Calculated", "계산값"),
    "proposal": ("Development proposal", "개발 제안"), "example": ("Fictional example; unvalidated; sources pending", "가상 예시; 미검증; 출처 확인 대기"),
    "api": ("API / drug substance", "API / 원료의약품"), "excipient": ("Excipient", "첨가제"), "starting_material": ("Starting material", "출발물질"), "finished_product": ("Finished drug product", "완제의약품"),
    "not_applicable": ("Not applicable", "해당없음"), "applicable": ("Applicable; rationale required", "적용; 근거 필요"),
    "compendial": ("Compendial baseline", "공정서 기초안"), "manufacturer_source": ("Manufacturer / DMF", "제조원 / DMF"), "approved": ("Approved specification", "승인 규격"), "experimental": ("Experimental data", "실험 자료"),
    "minimal": ("Minimal", "최소 접근법"), "enhanced": ("Enhanced (optional)", "강화된 접근법 (선택)"),
    "direct_analyte": ("Analyte material mass", "분석대상 원료 질량"), "product_mass": ("Whole product powder mass → nominal API", "완제 분말 질량 → 명목 API"), "product_volume": ("Whole product volume → nominal API", "완제 부피 → 명목 API"), "whole_matrix": ("Whole matrix mass (e.g. digestion)", "전체 매트릭스 질량 (예: 분해)"),
    "none": ("No potency correction (explicit basis)", "역가 보정 없음 (명시적 기준)"), "as_is": ("Assigned as-is potency (water already covered)", "지정 그대로 역가 (수분 이미 반영)"), "dry": ("Dry-basis purity × (1 − water fraction)", "건조물 순도 × (1 − 수분 분율)"),
    "sample": ("Sample solution", "검액"), "reference": ("Reference solution", "표준액"),
    "dmf_pending": ("DMF unreviewed — compendial baseline may proceed", "DMF 미검토 — 공정서 기초안 진행 가능"),
    "supplier_specification": ("Supplier specification", "제조원 규격"), "supplier_procedure": ("Supplier procedure", "제조원 분석절차"), "supplier_preparation": ("Supplier preparation", "제조원 조제법"), "supplier_impurities": ("Supplier impurity profile / gaps", "제조원 불순물 프로파일 / 차이"), "comparison_action": ("Gap / impact / next action", "차이 / 영향 / 다음 조치"),
    "limit_kind": ("Limit type / compound", "한도 종류 / 물질"), "allowance": ("Source-specific exposure allowance (ug/day)", "출처별 허용 노출량 (ug/day)"), "denominator": ("Daily mass denominator", "일일 질량 분모"), "whole_product": ("Whole drug product daily mass", "완제 전체 일일 질량"), "api_daily": ("API daily mass", "API 일일 질량"), "material_daily": ("Material daily contribution", "원료 일일 기여량"),
    "API dose differs from strength × units/day": ("API dose differs from strength × units/day", "API 투여량과 표시량 × 일일 투여단위수가 다릅니다"),
    "API strength exceeds whole product mass": ("API strength exceeds whole product mass", "API 표시량이 전체 제품 질량보다 큽니다"),
    "Reference requires analyte material mass": ("Reference requires analyte material mass", "표준액에는 분석대상 원료 질량이 필요합니다"),
    "Product-volume basis requires a liquid product and volume units": ("Product-volume basis requires a liquid product and volume units", "완제 부피 기준에는 액상 제품 및 부피 단위가 필요합니다"),
    "Product-mass basis requires a finished product with unit mass": ("Product-mass basis requires a finished product with unit mass", "완제 질량 기준에는 완제 구분 및 투여단위 질량이 필요합니다"),
    "Dry-basis purity must be ≤100%": ("Dry-basis purity must be ≤100%", "건조물 순도는 100% 이하여야 합니다"),
    "Water must be <100%": ("Water must be <100%", "수분은 100% 미만이어야 합니다"),
    "Extraction fraction must be ≤100%": ("Extraction fraction must be ≤100%", "추출 분율은 100% 이하여야 합니다"),
    "Whole-matrix mass must not receive analyte potency/extraction corrections": ("Whole-matrix mass must not receive analyte potency/extraction corrections", "전체 매트릭스 질량에 분석대상 역가 / 추출 보정을 적용하지 마세요"),
    "At most 8 explicit dilution steps": ("At most 8 explicit dilution steps", "명시적 희석 단계는 최대 8개입니다"),
    "Aliquot exceeds available or final volume": ("Aliquot exceeds available or final volume", "분취량이 원액량 또는 최종부피를 초과합니다"),
    "Nonfinite or underflowed concentration": ("Nonfinite or underflowed concentration", "유한한 양의 농도를 계산할 수 없습니다"),
    "Exposure allowance needs source, version, location and applicability": ("Exposure allowance needs source, version, location and applicability", "허용 노출량에는 문서·버전·위치·적용성 근거가 필요합니다"),
    "Conflicting daily API dose inputs": ("Conflicting daily API dose inputs", "일일 API 투여량 입력이 일치하지 않습니다"),
    "Whole product denominator requires finished-product inputs": ("Whole product denominator requires finished-product inputs", "전체 제품 분모에는 완제 입력이 필요합니다"),
    "Daily mass denominator unavailable": ("Daily mass denominator unavailable", "일일 질량 분모를 확인할 수 없습니다"),
    "J conversion requires whole-product daily mass and whole-matrix sample preparation": ("J conversion requires whole-product daily mass and whole-matrix sample preparation", "J 환산에는 전체 제품 일일 질량 및 전체 매트릭스 조제가 필요합니다"),
    "Enter a custom test name (1–160 characters)": ("Enter a custom test name (1–160 characters)", "사용자 시험항목명을 1–160자로 입력하세요"),
}

def label(key: str, lang: str = "en") -> str:
    return LABELS.get(key, (key, key))[1 if lang == "ko" else 0]

def test_label(key: str, state: dict[str, Any], lang: str = "en") -> str:
    if key in CATALOG:
        return CATALOG[key][1 if lang == "ko" else 0]
    return state["tests"][key].get("custom_name", key)

def blank_preparation(reference: bool = False) -> dict[str, Any]:
    return {"amount": None, "amount_unit": "mg", "basis": "direct_analyte", "correction_mode": "as_is" if reference else "none", "potency": None, "water": None, "basis_factor": 1.0, "basis_note": "", "extraction": 100.0, "stock_ml": None, "steps": [], "source": "", "version": "", "location": "", "value_kind": "entered", "review": None}

def blank_test(key: str) -> dict[str, Any]:
    fields = ("purpose", "rationale", "source", "version", "location", "scope", "excerpt", "specification", "procedure", "preparation", "sst", "verification_plan", "validation_plan", "experimental_data", "supplier_specification", "supplier_procedure", "supplier_preparation", "supplier_impurities", "comparison_action")
    result = {field: "" for field in fields}
    result.update({"applicability": "pending", "source_type": "compendial", "approach": "minimal", "value_kind": "proposal", "sample": blank_preparation(), "reference": blank_preparation(True), "evidence_review": None, "experiment_review": None})
    return result

def new_workspace() -> dict[str, Any]:
    product = {key: "" for key in ("name", "api_name", "identifier", "form", "claim_basis", "dosage_form", "route", "market", "manufacturer", "batch", "stage")}
    product.update({"material_type": "api", "max_daily_api": None, "api_dose_unit": "mg", "strength": None, "strength_unit": "mg", "whole_unit": None, "whole_unit_unit": "mg", "units_day": None, "density": None, "max_daily_material": None, "material_dose_unit": "mg"})
    return {"schema": 1, "build": BUILD, "product": product, "product_sources": [], "selected": [], "tests": {}, "dmf": {"source": "", "version": "", "location": "", "scope": ""}, "limit": {"limit_kind": "", "allowance": None, "denominator": "whole_product", "source": "", "version": "", "location": "", "scope": ""}, "example": False, "history": [], "signatures": {}}

def select_tests(state: dict[str, Any], keys: list[str]) -> None:
    if len(set(keys)) != len(keys) or any(key not in CATALOG and key not in state["tests"] for key in keys):
        raise ValueError("Unknown or duplicate test selection")
    state["selected"] = list(keys)
    for key in keys:
        state["tests"].setdefault(key, blank_test(key))
    for field in ("active_method_test", "active_calculation_test"):
        if state.get(field) not in keys:
            state[field] = keys[0] if keys else None

def add_custom_test(state: dict[str, Any], name: str) -> str:
    name = name.strip()
    if not name or len(name) > 160:
        raise ValueError("Enter a custom test name (1–160 characters)")
    key = "custom_" + sha256(name.casefold().encode()).hexdigest()[:16]
    if key not in state["tests"]:
        state["tests"][key] = blank_test(key)
        state["tests"][key]["custom_name"] = name
    select_tests(state, list(dict.fromkeys([*state["selected"], key])))
    return key

def positive(value: Any, field: str, zero: bool = False) -> float:
    if value is None or isinstance(value, bool):
        raise ValueError(field)
    number = float(value)
    if not isfinite(number) or number < 0 or (number == 0 and not zero):
        raise ValueError(field)
    return number

def mass_mg(value: Any, unit: str) -> float:
    if unit not in MASS_MG:
        raise ValueError("amount_unit")
    return positive(value, "amount") * MASS_MG[unit]

def daily_amounts(product: dict[str, Any]) -> dict[str, Any]:
    result: dict[str, Any] = {"api_mg_day": None, "implied_api_mg_day": None, "whole_g_day": None, "whole_ml_day": None, "material_g_day": None, "issues": []}
    for field, unit_field, output, divisor in (("max_daily_api", "api_dose_unit", "api_mg_day", 1), ("max_daily_material", "material_dose_unit", "material_g_day", 1000)):
        if product.get(field) is not None:
            try:
                result[output] = mass_mg(product[field], product[unit_field]) / divisor
            except ValueError:
                result["issues"].append(field)
    if product["material_type"] != "finished_product":
        return result
    try:
        units = positive(product.get("units_day"), "units_day")
    except ValueError:
        result["issues"].append("units_day")
        return result
    if product.get("strength") is not None:
        try:
            result["implied_api_mg_day"] = mass_mg(product["strength"], product["strength_unit"]) * units
            if result["api_mg_day"] is not None and not isclose(result["api_mg_day"], result["implied_api_mg_day"], rel_tol=1e-9, abs_tol=1e-9):
                result["issues"].append("API dose differs from strength × units/day")
        except ValueError:
            result["issues"].append("strength")
    try:
        whole = positive(product.get("whole_unit"), "whole_unit")
        unit = product["whole_unit_unit"]
        if unit in MASS_MG:
            result["whole_g_day"] = whole * MASS_MG[unit] * units / 1000
        elif unit in VOLUME_ML:
            result["whole_ml_day"] = whole * VOLUME_ML[unit] * units
            if product.get("density") is not None:
                result["whole_g_day"] = result["whole_ml_day"] * positive(product["density"], "density")
            else:
                result["issues"].append("density")
        else:
            result["issues"].append("whole_unit_unit")
        if unit in MASS_MG and product.get("strength") is not None and mass_mg(product["strength"], product["strength_unit"]) > whole * MASS_MG[unit]:
            result["issues"].append("API strength exceeds whole product mass")
    except ValueError as exc:
        result["issues"].append(str(exc))
    return result

def calculate_preparation(prep: dict[str, Any], product: dict[str, Any], reference: bool = False) -> dict[str, Any]:
    """Return nominal analyte concentration. Extraction is an explicit assumption."""
    basis = prep["basis"]
    if reference and basis != "direct_analyte":
        raise ValueError("Reference requires analyte material mass")
    if basis == "product_volume":
        if product["material_type"] != "finished_product" or prep["amount_unit"] not in VOLUME_ML or product["whole_unit_unit"] not in VOLUME_ML:
            raise ValueError("Product-volume basis requires a liquid product and volume units")
        original_mg = None
        amount_ml = positive(prep.get("amount"), "amount") * VOLUME_ML[prep["amount_unit"]]
        nominal_mg = amount_ml * mass_mg(product.get("strength"), product["strength_unit"]) / (positive(product.get("whole_unit"), "whole_unit") * VOLUME_ML[product["whole_unit_unit"]])
        amount_equation = f"{prep['amount']:g} {prep['amount_unit']} × {VOLUME_ML[prep['amount_unit']]:g} mL/{prep['amount_unit']} × {mass_mg(product['strength'], product['strength_unit']):g} mg API / {positive(product['whole_unit'], 'whole_unit') * VOLUME_ML[product['whole_unit_unit']]:g} mL product = {nominal_mg:g} mg API"
    else:
        original_mg = mass_mg(prep.get("amount"), prep["amount_unit"])
        nominal_mg = original_mg
        amount_equation = f"{prep['amount']:g} {prep['amount_unit']} × {MASS_MG[prep['amount_unit']]:g} mg/{prep['amount_unit']} = {original_mg:g} mg"
        if basis == "product_mass":
            if product["material_type"] != "finished_product" or product["whole_unit_unit"] not in MASS_MG:
                raise ValueError("Product-mass basis requires a finished product with unit mass")
            whole_mg = mass_mg(product.get("whole_unit"), product["whole_unit_unit"])
            strength_mg = mass_mg(product.get("strength"), product["strength_unit"])
            if strength_mg > whole_mg:
                raise ValueError("API strength exceeds whole product mass")
            nominal_mg *= strength_mg / whole_mg
            amount_equation += f" × ({strength_mg:g} mg API / {whole_mg:g} mg product) = {nominal_mg:g} mg API"
        elif basis not in {"direct_analyte", "whole_matrix"}:
            raise ValueError("basis")
    mode = prep["correction_mode"]
    if mode == "none":
        correction, correction_equation = 1.0, "f = 1 (explicit no-correction basis)"
    elif mode in {"as_is", "dry"}:
        potency = positive(prep.get("potency"), "potency")
        correction = potency / 100
        correction_equation = f"f = {potency:g}/100 = {correction:g}"
        if mode == "dry":
            if potency > 100:
                raise ValueError("Dry-basis purity must be ≤100%")
            water = positive(prep.get("water"), "water", zero=True)
            if water >= 100:
                raise ValueError("Water must be <100%")
            correction *= 1 - water / 100
            correction_equation = f"f = ({potency:g}/100) × (1 − {water:g}/100) = {correction:g}"
    else:
        raise ValueError("correction_mode")
    basis_factor = positive(prep.get("basis_factor"), "basis_factor")
    extraction = 1.0 if reference else positive(prep.get("extraction"), "extraction") / 100
    if extraction > 1:
        raise ValueError("Extraction fraction must be ≤100%")
    if basis == "whole_matrix" and (mode != "none" or basis_factor != 1 or extraction != 1):
        raise ValueError("Whole-matrix mass must not receive analyte potency/extraction corrections")
    corrected_mg = nominal_mg * correction * basis_factor * extraction
    stock = positive(prep.get("stock_ml"), "stock_ml")
    stock_ug_ml = corrected_mg * 1000 / stock
    current, previous_volume = stock_ug_ml, stock
    trace = [amount_equation, correction_equation, f"m_corrected = {nominal_mg:g} mg × {correction:g} × {basis_factor:g} × {extraction:g} = {corrected_mg:g} mg", f"C_stock = {corrected_mg:g} mg × 1000 ug/mg / {stock:g} mL = {stock_ug_ml:g} ug/mL"]
    steps = []
    if len(prep.get("steps", [])) > 8:
        raise ValueError("At most 8 explicit dilution steps")
    for index, step in enumerate(prep.get("steps", []), 1):
        aliquot = positive(step.get("aliquot_ml"), "aliquot_ml")
        final = positive(step.get("final_ml"), "final_ml")
        if aliquot > previous_volume or aliquot > final:
            raise ValueError("Aliquot exceeds available or final volume")
        before = current
        current *= aliquot / final
        equation = f"C_{index} = {before:g} ug/mL × ({aliquot:g} mL / {final:g} mL) = {current:g} ug/mL"
        steps.append({"step": index, "aliquot_ml": aliquot, "final_ml": final, "factor": final / aliquot, "concentration_ug_ml": current, "equation": equation})
        trace.append(equation)
        previous_volume = final
    values = (corrected_mg, stock_ug_ml, current)
    if any(not isfinite(value) or value <= 0 for value in values):
        raise ValueError("Nonfinite or underflowed concentration")
    return {"original_mg": original_mg, "nominal_mg": nominal_mg, "correction": correction, "basis_factor": basis_factor, "extraction": extraction, "corrected_mg": corrected_mg, "stock_ug_ml": stock_ug_ml, "final_ug_ml": current, "matrix_g_ml": current / 1_000_000 if basis == "whole_matrix" else None, "steps": steps, "trace": trace, "formula": "C_final [ug/mL] = nominal analyte mass [mg] × 1000 [ug/mg] × potency factor × analyte conversion factor × extraction fraction / V_stock [mL] × Π(V_aliquot [mL]/V_final [mL])"}

def impurity_limit(state: dict[str, Any], matrix_prep: dict[str, Any] | None = None) -> dict[str, Any]:
    limit = state["limit"]
    if any(not str(limit.get(field, "")).strip() for field in ("limit_kind", "source", "version", "location", "scope")):
        raise ValueError("Exposure allowance needs source, version, location and applicability")
    allowance = positive(limit.get("allowance"), "allowance")
    daily = daily_amounts(state["product"])
    basis = limit["denominator"]
    if basis == "api_daily":
        if "API dose differs from strength × units/day" in daily["issues"]:
            raise ValueError("Conflicting daily API dose inputs")
        grams = None if daily["api_mg_day"] is None else daily["api_mg_day"] / 1000
    elif basis == "whole_product":
        if state["product"]["material_type"] != "finished_product":
            raise ValueError("Whole product denominator requires finished-product inputs")
        grams = daily["whole_g_day"]
    elif basis == "material_daily":
        grams = daily["material_g_day"]
    else:
        raise ValueError("denominator")
    grams = positive(grams, "Daily mass denominator unavailable")
    value = allowance / grams
    result = {"ug_g": value, "g_day": grams, "equation": f"{allowance:g} ug/day ÷ {grams:g} g/day = {value:g} ug/g", "solution_ug_ml": None}
    if matrix_prep is not None:
        calculated = calculate_preparation(matrix_prep, state["product"])
        if matrix_prep["basis"] != "whole_matrix" or basis != "whole_product":
            raise ValueError("J conversion requires whole-product daily mass and whole-matrix sample preparation")
        result["solution_ug_ml"] = value * calculated["matrix_g_ml"]
        result["equation"] += f"; J = {value:g} ug/g × {calculated['matrix_g_ml']:g} g/mL = {result['solution_ug_ml']:g} ug/mL"
    return result

def _without_reviews(value: Any) -> Any:
    if isinstance(value, dict):
        return {k: _without_reviews(v) for k, v in value.items() if k not in {"review", "evidence_review", "experiment_review"}}
    if isinstance(value, list):
        return [_without_reviews(v) for v in value]
    return value

def fingerprint(value: Any) -> str:
    return sha256(json.dumps(_without_reviews(value), sort_keys=True, ensure_ascii=False, allow_nan=False).encode()).hexdigest()

def review_payload(state: dict[str, Any], key: str, kind: str) -> Any:
    record = state["tests"][key]
    common = [BUILD, state["product"], state["product_sources"], state["dmf"], state.get("context", {}), state.get("example", False)]
    metadata = {field: v for field, v in record.items() if field not in {"sample", "reference", "evidence_review", "experiment_review", "experimental_data"}}
    return [*common, record if kind in {"evidence", "experiment"} else [metadata, record[kind]]]

def refresh_reviews(state: dict[str, Any], context: dict[str, Any] | None = None) -> None:
    if context is not None:
        state["context"] = dict(context)
    for key, record in state["tests"].items():
        for kind in ("sample", "reference", "evidence", "experiment"):
            signature = fingerprint(review_payload(state, key, kind))
            slot = key + ":" + kind
            old = state["signatures"].get(slot)
            holder, field = (record[kind], "review") if kind in {"sample", "reference"} else (record, kind + "_review")
            if old is not None and old != signature and holder.get(field):
                state["history"].append({"test": key, "kind": kind, "reason": "Inputs / source / scope changed", "prior_review": holder[field]})
                holder[field] = None
            # A stale saved review cannot revive by restoring old values.
            if holder.get(field) and holder[field].get("fingerprint") != signature:
                holder[field] = None
            state["signatures"][slot] = signature
    state["history"] = state["history"][-100:]

def record_review(state: dict[str, Any], key: str, kind: str, reviewer: str, note: str) -> None:
    refresh_reviews(state)
    if not reviewer.strip() or not note.strip():
        raise ValueError("reviewer / review_note")
    record = state["tests"][key]
    if kind in {"sample", "reference"}:
        calculate_preparation(record[kind], state["product"], reference=kind == "reference")
        holder, field = record[kind], "review"
    elif kind in {"evidence", "experiment"}:
        required = ["source", "version", "location", "scope", "rationale"]
        if record["applicability"] == "pending":
            raise ValueError("applicability")
        if kind == "experiment":
            required += ["validation_plan", "experimental_data"]
        elif record["applicability"] == "applicable":
            required += ["excerpt", "specification", "procedure", "preparation", "sst", "verification_plan"]
        missing = [field for field in required if not str(record.get(field, "")).strip()]
        if missing:
            raise ValueError(", ".join(missing))
        holder, field = record, kind + "_review"
    else:
        raise ValueError("review kind")
    holder[field] = {"reviewer": reviewer.strip(), "note": note.strip(), "fingerprint": fingerprint(review_payload(state, key, kind)), "at": datetime.now(timezone.utc).isoformat()}

def comparison_rows(record: dict[str, Any], lang: str = "en") -> list[list[str]]:
    rows = []
    for field in ("specification", "procedure", "preparation"):
        baseline, supplier = record.get(field, ""), record.get("supplier_" + field, "")
        status = ("Comparison pending", "대조 대기") if not baseline.strip() or not supplier.strip() else (("Text matches; equivalence unreviewed", "문구 일치; 동등성 미검토") if baseline.strip() == supplier.strip() else ("Difference; impact review needed", "차이; 영향 검토 필요"))
        rows.append([label(field, lang), baseline, supplier, status[lang == "ko"]])
    return rows

def demo_workspace() -> dict[str, Any]:
    state = new_workspace()
    state["example"] = True
    state["product"].update({"material_type": "finished_product", "name": "Fictional demonstration tablet", "api_name": "Fictional API", "claim_basis": "Nominal API label claim; extraction assumed 100%", "dosage_form": "Tablet", "max_daily_api": 400.0, "strength": 100.0, "whole_unit": 300.0, "units_day": 4.0})
    select_tests(state, ["identity", "assay", "related_substances", "dissolution"])
    record = state["tests"]["assay"]
    record["purpose"] = "Fictional assay calculation demonstration; not a validated procedure"
    record["sample"].update({"amount": 300.0, "basis": "product_mass", "stock_ml": 100.0, "steps": [{"aliquot_ml": 5.0, "final_ml": 50.0}], "value_kind": "example", "basis_note": "Homogeneous powder, label claim and 100% extraction assumed; sources pending"})
    record["reference"].update({"amount": 20.0, "potency": 99.5, "stock_ml": 100.0, "steps": [{"aliquot_ml": 5.0, "final_ml": 10.0}], "value_kind": "example", "basis_note": "Fictional assigned as-is potency; actual CoA verification pending"})
    return state

def _safe_text(value: Any) -> str:
    # Export text stays text even when users enter spreadsheet formula prefixes.
    return str(value if value is not None else "").replace("\x00", "")

def report_markdown(state: dict[str, Any], lang: str = "en") -> str:
    refresh_reviews(state)
    ko = lang == "ko"
    def title(en: str, kr: str) -> str:
        return kr if ko else en
    def value(v: Any) -> str:
        if v is None or v == "":
            return title("Not recorded", "미기록")
        return label(v, lang) if isinstance(v, str) and v in LABELS else str(v)
    def line(k: str, v: Any) -> str:
        return f"- {label(k, lang)}: {value(v)}\n"
    report = f"# {title('Analytical method workspace', '분석법 작업공간')}\n\n{BUILD}\n\n"
    report += title("Input completion, source verification and experimental support are separate. Review records are session declarations, not QC use approval. No universal concentration or pharmacopeial threshold is supplied.\n\n", "입력 완료, 근거 확인과 실험적 뒷받침은 별도입니다. 검토 기록은 세션의 선언이며 QC 사용 승인이 아닙니다. 보편적 농도나 공정서 한도를 제공하지 않습니다.\n\n")
    if state["example"]:
        report += label("example", lang) + "\n\n"
    report += "## " + title("Product inputs", "제품 입력") + "\n\n"
    for k, v in state["product"].items():
        # Fields inactive for raw materials are retained but are not dose denominators.
        report += line(k, v)
    daily = daily_amounts(state["product"])
    report += title("\nCalculated daily amounts", "\n계산된 일일량") + f": API {daily['api_mg_day']} mg/day; strength × units/day {daily['implied_api_mg_day']} mg/day; whole product {daily['whole_g_day']} g/day; volume {daily['whole_ml_day']} mL/day.\n\n"
    if daily["issues"]:
        report += title("Missing / conflicting dose basis: ", "투여량 근거 누락 / 불일치: ") + ", ".join(label(item, lang) for item in daily["issues"]) + "\n\n"
    report += title("### Per-value provenance\n\n", "### 값별 출처\n\n")
    for source in state["product_sources"]:
        report += "; ".join(f"{label(k, lang)}: {value(v)}" for k, v in source.items()) + "\n\n"
    report += label("dmf_pending", lang) + "\n\n"
    for field, v in state["dmf"].items():
        report += line(field, v)
    report += "## " + title("Selected tests", "선택 시험") + "\n\n"
    for key in state["selected"]:
        record = state["tests"][key]
        report += "### " + test_label(key, state, lang) + "\n\n"
        for field, v in record.items():
            if field in {"sample", "reference", "evidence_review", "experiment_review", "input_sources"}:
                continue
            report += line(field, v)
        for origin in record.get("input_sources", []):
            report += "; ".join(f"{label(k, lang)}: {value(v)}" for k, v in origin.items()) + "\n\n"
        for kind in ("evidence", "experiment"):
            review = record.get(kind + "_review")
            report += title("Source verification" if kind == "evidence" else "Experimental support", "근거 확인" if kind == "evidence" else "실험적 뒷받침") + ": " + label("reviewed" if review else "pending", lang) + "\n\n"
            if review:
                report += json.dumps(review, ensure_ascii=False) + "\n\n"
        for kind in ("sample", "reference"):
            prep = record[kind]
            report += "#### " + label(kind, lang) + "\n\n"
            for k, v in prep.items():
                if k not in {"review", "steps"}:
                    if k == "input_sources":
                        continue
                    report += line(k, v)
            for origin in prep.get("input_sources", []):
                report += "; ".join(f"{label(k, lang)}: {value(v)}" for k, v in origin.items()) + "\n\n"
            report += label("reviewed" if prep.get("review") else "pending", lang) + "\n\n"
            if prep.get("review"):
                report += json.dumps(prep["review"], ensure_ascii=False) + "\n\n"
            try:
                calculated = calculate_preparation(prep, state["product"], kind == "reference")
                report += calculated["formula"] + "\n\n" + "\n\n".join(calculated["trace"]) + "\n\n"
                report += f"{label('calculated', lang)}: {calculated['final_ug_ml']:.12g} ug/mL\n\n"
            except (ValueError, TypeError, OverflowError) as exc:
                report += label("incomplete", lang) + ": " + label(str(exc), lang) + "\n\n"
        for row in comparison_rows(record, lang):
            report += " | ".join(row) + "\n\n"
    report += "## " + title("Exposure allowance conversion", "허용 노출량 환산") + "\n\n"
    for field, v in state["limit"].items():
        report += line(field, v)
    try:
        report += impurity_limit(state)["equation"] + "\n\n"
    except ValueError as exc:
        report += label("pending", lang) + ": " + label(str(exc), lang) + "\n\n"
    report += title("### Review invalidation history\n\n", "### 검토 무효화 이력\n\n") + json.dumps(state["history"], ensure_ascii=False, indent=2) + "\n"
    return report

def export_workbook(state: dict[str, Any], lang: str = "en") -> bytes:
    """Editable raw inputs and cell-linked formulas; no cached result masquerades."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.workbook.properties import CalcProperties

    refresh_reviews(state)
    wb = Workbook()
    wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
    ws = wb.active
    ws.title = "Read me"
    rows = [[BUILD], [label("example", lang) if state["example"] else label("proposal", lang)],
            ["수식 셀은 Excel에서 다시 계산합니다. 검토 기록은 내보낸 입력 스냅샷에만 유효합니다. 셀 편집 후 재검토하세요." if lang == "ko" else "Formula cells recalculate in Excel. Review records apply only to the exported input snapshot; review again after editing cells."],
            ["청색 입력값 / 녹색 수식. 원문은 번역하지 않습니다." if lang == "ko" else "Blue inputs / green formulas. Source text is preserved."],
            ["sample concentration = analyte mass × correction × extraction / stock volume × serial dilution ratios"],
            ["Q14 minimal approach accepted; formal ATP / DoE / MODR are not universal requirements. Q2(R2) validation differs from Q6 specifications and Q3/M7 exposure limits."],
            ["No preset exposure allowance, official method or QC approval is generated."]]
    for row in rows:
        ws.append(row)
    product_ws = wb.create_sheet("Product")
    product_ws.append(["Field" if lang == "en" else "항목", "Value" if lang == "en" else "값", "Unit / basis" if lang == "en" else "단위 / 기준"])
    product_cells = {}
    product_origins = {row["field"]: row.get("value_kind", "entered") for row in state["product_sources"]}
    for field, v in state["product"].items():
        row = product_ws.max_row + 1
        product_ws.append([label(field, lang), v, label(product_origins.get(field, "entered"), lang)])
        product_cells[field] = f"'Product'!B{row}"
    def pc(field: str) -> str:
        return product_cells[field]
    def factor(cell: str, mapping: dict[str, float]) -> str:
        expression = "NA()"
        for unit, multiplier in reversed(list(mapping.items())):
            expression = f'IF({cell}="{unit}",{multiplier},{expression})'
        return expression
    provenance = wb.create_sheet("Provenance")
    provenance.append(["Input field", label("value_kind", lang), label("source", lang), label("version", lang), label("location", lang), label("basis_note", lang)])
    for source in state["product_sources"]:
        provenance.append([source.get(k, "") for k in ("field", "value_kind", "source", "version", "location", "basis_note")])
    evidence = wb.create_sheet("Evidence")
    evidence.append(["Test", "Field", "Value"] if lang == "en" else ["시험", "항목", "값"])
    formulas: set[tuple[str, str]] = set()
    for idx, key in enumerate(state["selected"], 1):
        record = state["tests"][key]
        for origin in record.get("input_sources", []):
            provenance.append([f"{test_label(key, state, lang)} / {label(origin.get('field', ''), lang)}", *[origin.get(k, "") for k in ("value_kind", "source", "version", "location", "basis_note")]])
        for field, v in record.items():
            if field not in {"sample", "reference", "input_sources"}:
                evidence.append([test_label(key, state, lang), label(field, lang), json.dumps(v, ensure_ascii=False) if isinstance(v, dict) else v])
    for field, v in state["dmf"].items():
        evidence.append([label("dmf_pending", lang), label(field, lang), v])
    for idx, key in enumerate(state["selected"], 1):
        record = state["tests"][key]
        for kind in ("sample", "reference"):
            prep = record[kind]
            for source in prep.get("input_sources", []):
                provenance.append([f"{test_label(key, state, lang)} / {label(kind, lang)} / {label(source.get('field', ''), lang)}", *[source.get(k, "") for k in ("value_kind", "source", "version", "location", "basis_note")]])
            sheet = wb.create_sheet(f"{idx:02d}_{kind}")
            sheet.append([test_label(key, state, lang), label(kind, lang), "Units", "Basis"])
            inputs = [prep["amount"], MASS_MG.get(prep["amount_unit"], VOLUME_ML.get(prep["amount_unit"])), None, None, prep.get("potency"), prep.get("water"), None, prep.get("basis_factor"), 100.0 if kind == "reference" else prep.get("extraction"), prep.get("stock_ml")]
            keys = ["amount", "amount_unit", "strength", "whole_unit", "potency", "water", "correction_mode", "basis_factor", "extraction", "stock_ml"]
            for row, (field, v) in enumerate(zip(keys, inputs), 2):
                sheet.append([label(field, lang), v, prep["amount_unit"] if row == 2 else "", label(prep.get("basis"), lang) if row == 2 else ""])
            unit_mapping = VOLUME_ML if prep["basis"] == "product_volume" else MASS_MG
            sheet["B3"] = "=" + factor("C2", unit_mapping)
            sheet["B4"] = f"={pc('strength')}*{factor(pc('strength_unit'), MASS_MG)}" if prep["basis"] in {"product_mass", "product_volume"} else 1
            whole_factor = MASS_MG.get(state["product"]["whole_unit_unit"], VOLUME_ML.get(state["product"]["whole_unit_unit"]))
            sheet["B5"] = f"={pc('whole_unit')}*{factor(pc('whole_unit_unit'), unit_mapping)}" if prep["basis"] in {"product_mass", "product_volume"} else 1
            sheet["B8"] = "=B6/100*(1-B7/100)" if prep["correction_mode"] == "dry" else ("=B6/100" if prep["correction_mode"] == "as_is" else "=1")
            sheet.append(["Nominal mass (mg)" if lang == "en" else "명목 질량 (mg)", "=B2*B3*B4/B5", "mg"])
            sheet.append(["Corrected mass (mg)" if lang == "en" else "보정 질량 (mg)", "=B12*B8*B9*B10/100", "mg"])
            sheet.append(["Stock concentration" if lang == "en" else "원액 농도", "=B13*1000/B11", "ug/mL"])
            previous = "B14"
            sheet.append(["Step" if lang == "en" else "단계", label("aliquot_ml", lang), label("final_ml", lang), "C (ug/mL)"])
            for row, step in enumerate(prep["steps"], 16):
                sheet.append([row - 15, step.get("aliquot_ml"), step.get("final_ml"), f"={previous}*B{row}/C{row}"])
                previous = f"D{row}"
            row = sheet.max_row + 1
            sheet.append(["Final concentration" if lang == "en" else "최종 농도", "=" + previous, "ug/mL"])
            try:
                calc = calculate_preparation(prep, state["product"], kind == "reference")
                sheet.append(["App calculation snapshot" if lang == "en" else "앱 계산 스냅샷", calc["final_ug_ml"], "ug/mL"])
                sheet.append(["Compare to snapshot (ug/mL)" if lang == "en" else "스냅샷 차이 (ug/mL)", f"=B{row}-B{row+1}"])
                for equation in calc["trace"]:
                    sheet.append([equation])
            except (ValueError, TypeError, OverflowError) as exc:
                # Incomplete preparations contain inputs only, no working-looking formulas.
                for cells in sheet:
                    for cell in cells:
                        if cell.data_type == "f":
                            cell.value = None
                sheet.append([label("incomplete", lang), label(str(exc), lang)])
            for field in ("source", "version", "location", "basis_note", "value_kind"):
                sheet.append([label(field, lang), prep.get(field, "")])
            sheet.append([label("reviewed" if prep.get("review") else "pending", lang), json.dumps(prep.get("review"), ensure_ascii=False)])
            for cells in sheet:
                for cell in cells:
                    if cell.data_type == "f":
                        formulas.add((sheet.title, cell.coordinate))
    limit_ws = wb.create_sheet("Exposure")
    for field, v in state["limit"].items():
        limit_ws.append([label(field, lang), v])
    try:
        limit = impurity_limit(state)
        # Both API dose and whole-product denominators are cell-linked, never interchangeable.
        denominator = state["limit"]["denominator"]
        if denominator == "api_daily":
            denominator_formula = f"={pc('max_daily_api')}*{factor(pc('api_dose_unit'), MASS_MG)}/1000"
        elif denominator == "material_daily":
            denominator_formula = f"={pc('max_daily_material')}*{factor(pc('material_dose_unit'), MASS_MG)}/1000"
        else:
            unit = state["product"]["whole_unit_unit"]
            denominator_formula = f"={pc('whole_unit')}*{pc('units_day')}*{factor(pc('whole_unit_unit'), MASS_MG)}/1000" if unit in MASS_MG else f"={pc('whole_unit')}*{factor(pc('whole_unit_unit'), VOLUME_ML)}*{pc('units_day')}*{pc('density')}"
        row = limit_ws.max_row + 1
        limit_ws.append([label("denominator", lang) + " (g/day)", denominator_formula])
        limit_ws.append(["ug/g", f"=B2/B{row}"])
        limit_ws.append(["App snapshot (ug/g)", limit["ug_g"]])
        formulas.update({(limit_ws.title, f"B{row}"), (limit_ws.title, f"B{row+1}")})
    except ValueError as exc:
        limit_ws.append([label("pending", lang), label(str(exc), lang)])
    for sheet in wb:
        sheet.freeze_panes = "A2"
        for cells in sheet:
            for cell in cells:
                is_formula = (sheet.title, cell.coordinate) in formulas
                if isinstance(cell.value, str) and not is_formula:
                    cell.value = _safe_text(cell.value)
                    cell.data_type = "s"
                cell.font = Font(color="087F5B" if is_formula else "2456A6" if cell.column > 1 else "163945")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="E7F2F1")
            cell.font = Font(bold=True, color="163945")
        for col, width in (("A", 48), ("B", 44), ("C", 24), ("D", 48), ("E", 30), ("F", 48)):
            sheet.column_dimensions[col].width = width
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
